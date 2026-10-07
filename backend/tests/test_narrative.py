import io
import json
import uuid

import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pydantic import ValidationError

from app.api.deps import get_queue
from app.domain.brief import PresentationBrief, brief_directive
from app.domain.narrative import PROFILES, NarrativePlan, preset_plan, semantic_skeleton
from app.domain.narrative_quality import review_narrative
from app.domain.outline import OutlineDraft, OutlinePageDraft
from app.llm.base import OutlineGenerationInput
from app.llm.deepseek import DeepSeekOutlineGenerator
from app.llm.demo import DemoOutlineGenerator, DemoSlideGenerator
from app.llm.errors import InvalidModelOutputError
from app.main import app
from app.worker.deck_tasks import generate_deck
from app.worker.tasks import generate_outline
from app.workflows.outline import build_outline_workflow, run_outline_workflow


def brief(profile="executive", **kwargs):
    return PresentationBrief(narrative_enabled=True, audience_profile=profile, **kwargs)


def page(**kwargs):
    return OutlinePageDraft(
        title="可追溯的结论",
        objective="解释证据与限制",
        key_points=["已有材料", "待验证结论"],
        layout_id="bullets",
        **kwargs,
    )


@pytest.mark.parametrize("profile", list(PROFILES))
def test_each_profile_has_editable_needs_and_distinct_arc(profile):
    plan = preset_plan(brief(profile), "具体的听众")
    assert plan.audience_summary == "具体的听众"
    assert len(plan.arc) >= 3 and len(plan.audience_needs) >= 2
    assert plan.source == "preset"
    assert "设计假设" in "".join(plan.assumptions)
    assert len({tuple(preset_plan(brief(p)).arc) for p in PROFILES}) == len(PROFILES)


def test_legacy_brief_compatible_and_new_brief_avoids_academic_sequence():
    assert "按背景、问题、方案" in brief_directive(PresentationBrief(scenario="defense"))
    directive = brief_directive(brief(narrative_overrides="先给决策建议", knowledge_level="expert"))
    assert "先给决策建议" in directive and "expert" in directive
    assert "按背景、问题、方案" not in directive
    with pytest.raises(ValidationError):
        brief("unsupported")


@pytest.mark.parametrize("kind", ["auto", "claim", "comparison", "process", "data", "image"])
def test_semantic_layout_never_forces_metrics_without_data(kind):
    hint = semantic_skeleton(kind, has_data=False, position=3)
    assert "3 个 kpi" not in hint and "数据证据页" not in hint
    assert "数据证据页" in semantic_skeleton("data", has_data=True, position=1)


async def test_ai_plan_passes_to_outline_and_time_budget_is_preserved():
    calls = []
    expected = preset_plan(brief()).model_copy(
        update={"throughline": "先决定是否试点，再讲验证风险"}
    )

    class Chat:
        async def complete(self, schema, **kwargs):
            calls.append((schema, kwargs))
            return (
                expected
                if schema is NarrativePlan
                else OutlineDraft(pages=[page(page_role="cover"), page()])
            )

    payload = OutlineGenerationInput(
        title="试点提案", tone="professional", page_count=2, brief=brief()
    )
    draft = await run_outline_workflow(
        build_outline_workflow(DeepSeekOutlineGenerator(chat=Chat())), payload
    )
    assert draft.narrative.source == "ai"
    assert "先决定是否试点" in calls[1][1]["user"]
    assert (
        json.loads(calls[0][1]["user"])["principles"]["cognitive_load"]["citation"]
        == "Sweller (1988)"
    )
    assert sum(p.speaker_seconds for p in draft.pages) == 300
    assert draft.pages[0].speaker_seconds < draft.pages[1].speaker_seconds


@pytest.mark.parametrize("count", [5, 10, 15, 20])
@pytest.mark.parametrize("minutes", [1, 20])
async def test_page_count_matrix_preserves_time_budget(count, minutes):
    """Deterministic workflow coverage; separate from real-provider acceptance."""
    payload = OutlineGenerationInput(
        title="不同页数回归",
        tone="professional",
        page_count=count,
        brief=brief(duration_minutes=minutes),
    )
    draft = await run_outline_workflow(build_outline_workflow(DemoOutlineGenerator()), payload)
    assert len(draft.pages) == count
    assert all(p.speaker_seconds >= 1 for p in draft.pages)
    assert sum(p.speaker_seconds for p in draft.pages) == minutes * 60
    assert draft.narrative.source == "preset"


async def test_planning_failure_is_visible_and_outline_still_runs():
    class Chat:
        async def complete(self, schema, **kwargs):
            if schema is NarrativePlan:
                raise InvalidModelOutputError("Bad provider output")
            return OutlineDraft(pages=[page()])

    payload = OutlineGenerationInput(
        title="试点提案", tone="professional", page_count=1, brief=brief()
    )
    draft = await run_outline_workflow(
        build_outline_workflow(DeepSeekOutlineGenerator(chat=Chat())), payload
    )
    assert draft.narrative.source == "preset"
    assert "AI 规划未完成" in "".join(draft.narrative.assumptions)
    assert draft.pages[0].speaker_seconds == 300


def test_editorial_warnings_do_not_block_export():
    pages = [page(visual_kind="claim", speaker_seconds=10) for _ in range(4)]
    issues = review_narrative(pages, duration_minutes=5)
    assert {"narrative_timing", "narrative_repeated_title", "narrative_visual_rhythm"} <= {
        i.code for i in issues
    }
    assert all(i.severity == "warning" for i in issues)


async def test_narrative_persists_reaches_slide_generation_and_export_and_invalidates():
    class Queue:
        calls = []

        async def enqueue_job(self, *args, **kwargs):
            self.calls.append(args)
            return object()

    queue = Queue()
    received = []

    class Slides(DemoSlideGenerator):
        async def generate(self, payload):
            received.append(payload)
            return await super().generate(payload)

    app.dependency_overrides[get_queue] = lambda: queue
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            auth = await client.post(
                "/api/v1/auth/register",
                json={
                    "email": f"narrative-{uuid.uuid4().hex}@example.com",
                    "password": "test-password",
                },
            )
            client.headers["Authorization"] = f"Bearer {auth.json()['access_token']}"
            result = await client.post(
                "/api/v1/projects",
                json={"title": "听众验收", "page_count": 6, "brief": brief().model_dump()},
            )
            assert result.status_code == 201
            project = result.json()
            url = f"/api/v1/projects/{project['id']}"
            await client.post(
                url + "/sources",
                json={
                    "kind": "text",
                    "content": (
                        "问题：校园信息分散。\n\n方案：集中发布与搜索。\n\n边界：没有真实用户反馈。"
                    ),
                },
            )
            accepted = await client.post(url + "/outline/generate")
            await generate_outline(
                {"outline_generator": DemoOutlineGenerator()},
                project["id"],
                accepted.json()["job_id"],
            )
            outline = (await client.get(url + "/outline")).json()
            assert outline["status"] == "draft" and outline["narrative"]["source"] == "preset"
            assert outline["narrative"]["arc"][0] == "建议与待决事项"
            assert sum(p["speaker_seconds"] for p in outline["pages"]) == 300
            await client.patch(url, json={"brief": brief("learner").model_dump()})
            assert (
                await client.post(url + "/outline/confirm", json={"revision": outline["revision"]})
            ).status_code == 409
            await client.patch(url, json={"brief": brief().model_dump()})
            assert (
                await client.post(url + "/outline/confirm", json={"revision": outline["revision"]})
            ).status_code == 200
            assert (await client.post(url + "/deck/generate", json={})).status_code == 202
            _, project_id, slide_ids = queue.calls[-1]
            await generate_deck({"slide_generator": Slides()}, project_id, slide_ids)
            assert len(received) == 6 and all(
                p.narrative and p.narrative.arc[0] == "建议与待决事项" for p in received
            )
            assert all("3 个 kpi" not in (p.skeleton_hint or "") for p in received)
            exported = await client.get(url + "/deck/export")
            assert exported.status_code == 200, exported.text
            ppt = Presentation(io.BytesIO(exported.content))
            assert all(
                "讲述提示" in slide.notes_slide.notes_text_frame.text for slide in ppt.slides
            )
            assert all(
                "建议用时" in slide.notes_slide.notes_text_frame.text for slide in ppt.slides
            )
            await client.delete(url)
    finally:
        app.dependency_overrides.pop(get_queue, None)


def test_novel_numeric_targets_are_removed_from_planning():
    from app.domain.narrative import guard_plan_numbers

    baseline = preset_plan(brief())
    proposal = baseline.model_copy(
        update={
            "arc": ["12条样本的验证", "两周试点，成功率80%", "申请预算5000元"],
            "source": "ai",
        }
    )
    guarded = guard_plan_numbers(proposal, baseline, "本地验证12条记录；8条待认领4条已认领")
    assert guarded.arc == baseline.arc
    assert "数值" in guarded.assumptions[-1]
    assert guarded.source == "ai"


def test_optional_llm_nulls_preserve_missing_evidence():
    from app.domain.evidence import PointEvidence, sanitize_evidence

    p = page(
        transition=None,
        narrative_role=None,
        visual_kind=None,
        point_evidence=[PointEvidence(point="已有材料", kind="missing", quote=None)],
    )
    assert p.transition == "" and p.visual_kind == "auto"
    assert sanitize_evidence(p, {}).point_evidence[0].kind == "missing"


@pytest.mark.parametrize("claim", ["先开展两周试点", "计划2-3周完成", "申请1人支持"])
def test_small_or_chinese_quantity_not_silently_introduced_in_plan(claim):
    from app.domain.narrative import guard_plan_numbers
    from app.domain.quality import extract_numbers

    baseline = preset_plan(brief())
    proposal = baseline.model_copy(update={"closing_action": claim})
    guarded = guard_plan_numbers(proposal, baseline, "12条记录；8条待认领4条已认领")
    assert guarded.closing_action == baseline.closing_action
    assert extract_numbers(claim, include_quantities=True)


def test_narrative_outline_prompt_exposes_constraints_to_json_mode_model():
    generator = DeepSeekOutlineGenerator(chat=object(), layout_ids=frozenset({"cover"}))
    prompt = generator._system_prompt(narrative=True)
    assert '"maxLength": 180' in prompt
    assert '"maxItems": 5' in prompt
