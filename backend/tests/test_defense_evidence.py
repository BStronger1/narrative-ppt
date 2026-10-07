import io
import uuid

import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation

from app.api.deps import get_queue
from app.domain.evidence import PointEvidence, sanitize_evidence
from app.domain.outline import OutlinePageDraft
from app.llm.base import OutlineSourceSection
from app.llm.demo import DemoOutlineGenerator, DemoSlideGenerator
from app.main import app
from app.worker.deck_tasks import generate_deck
from app.worker.tasks import generate_outline


def page(**overrides):
    return OutlinePageDraft(
        **{
            "title": "验证结果",
            "objective": "展示结果",
            "layout_id": "bullets",
            "key_points": ["已完成本地验证", "需要用户反馈"],
            "source_refs": ["S1:1"],
            "point_evidence": [
                PointEvidence(
                    point="已完成本地验证",
                    kind="source",
                    ref="S1:1",
                    quote="已完成本地验证",
                )
            ],
            **overrides,
        }
    )


@pytest.mark.parametrize("ref,quote", [("S9:1", "已完成本地验证"), ("S1:1", "提升80%")])
def test_invalid_evidence_becomes_missing(ref, quote):
    sections = {
        "S1:1": OutlineSourceSection(
            ref="S1:1",
            level=0,
            locator="第1段",
            text="项目已完成本地验证。",
        )
    }
    draft = page(
        point_evidence=[
            PointEvidence(
                point="已完成本地验证",
                kind="source",
                ref=ref,
                quote=quote,
            )
        ]
    )
    assert sanitize_evidence(draft, sections).point_evidence[0].kind == "missing"


def test_editing_claim_invalidates_old_evidence_and_topic_is_not_proof():
    sections = {
        "S1:1": OutlineSourceSection(
            ref="S1:1",
            level=0,
            locator="第1段",
            text="项目已完成本地验证。",
        )
    }
    assert sanitize_evidence(page(), sections).point_evidence[0].kind == "source"
    changed = page(key_points=["已完成真实用户验证", "需要用户反馈"])
    assert sanitize_evidence(changed, sections).point_evidence[0].kind == "missing"
    assert (
        sanitize_evidence(page(), sections, topic_refs={"S1:1"}).point_evidence[0].kind == "missing"
    )


@pytest.mark.asyncio
async def test_defense_end_to_end_export_and_brief_invalidation():
    class Queue:
        calls = []

        async def enqueue_job(self, *args, **kwargs):
            self.calls.append(args)
            return object()

    queue = Queue()
    app.dependency_overrides[get_queue] = lambda: queue
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            auth = await client.post(
                "/api/v1/auth/register",
                json={
                    "email": f"defense-{uuid.uuid4().hex}@example.com",
                    "password": "test-password",
                },
            )
            client.headers["Authorization"] = f"Bearer {auth.json()['access_token']}"
            response = await client.post(
                "/api/v1/projects",
                json={
                    "title": "答辩集成验证",
                    "page_count": 6,
                    "layout_mode": "flex",
                    "brief": {"scenario": "defense", "duration_minutes": 5, "focus": "我的贡献"},
                },
            )
            assert response.status_code == 201, response.text
            project = response.json()
            assert project["brief"]["focus"] == "我的贡献"
            url = f"/api/v1/projects/{project['id']}"
            result = await client.post(
                f"{url}/sources",
                json={
                    "kind": "text",
                    "content": "\n\n".join(
                        [
                            "背景：校园信息分散在多个群聊。查找需要人工翻阅消息。",
                            "问题：失物记录缺乏统一入口。需要提供按分类和关键词查找功能。",
                            "方案：统一管理发布记录。使用关键词检索帮助定位相关物品。",
                            "实现：前端采用 React。后端采用 FastAPI 与 PostgreSQL。",
                            "结果：已完成本地功能验证。目前没有真实用户测试数据。",
                            "不足：认领需要人工确认。下一步计划补充异常输入测试。",
                        ]
                    ),
                },
            )
            assert result.status_code == 201, result.text
            accepted = await client.post(f"{url}/outline/generate")
            assert accepted.status_code == 202, accepted.text
            await generate_outline(
                {"outline_generator": DemoOutlineGenerator()},
                project["id"],
                accepted.json()["job_id"],
            )
            outline = (await client.get(f"{url}/outline")).json()
            assert outline["pages"][0]["point_evidence"][0]["kind"] == "source"
            patch = await client.patch(
                url,
                json={
                    "brief": {
                        "scenario": "defense",
                        "duration_minutes": 10,
                        "focus": "我的贡献",
                    }
                },
            )
            assert patch.status_code == 200
            stale = await client.post(
                f"{url}/outline/confirm", json={"revision": outline["revision"]}
            )
            assert stale.status_code == 409
            await client.patch(url, json={"brief": project["brief"]})
            confirmed = await client.post(
                f"{url}/outline/confirm", json={"revision": outline["revision"]}
            )
            assert confirmed.status_code == 200, confirmed.text
            generated = await client.post(f"{url}/deck/generate", json={})
            assert generated.status_code == 202, generated.text
            _, project_id, slide_ids = queue.calls[-1]
            await generate_deck({"slide_generator": DemoSlideGenerator()}, project_id, slide_ids)
            exported = await client.get(f"{url}/deck/export")
            assert exported.status_code == 200, exported.text
            ppt = Presentation(io.BytesIO(exported.content))
            assert len(ppt.slides) == 6
            assert any(shape.has_text_frame for shape in ppt.slides[0].shapes)
            notes = ppt.slides[0].notes_slide.notes_text_frame.text
            assert "S1:1" in notes and "校园信息分散" in notes
            assert "演示模式" in notes
            await client.delete(url)
    finally:
        app.dependency_overrides.pop(get_queue, None)
