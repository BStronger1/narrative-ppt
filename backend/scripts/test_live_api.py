"""手动验收真实模型链路（会消耗已配置的 API 额度，不由 pytest 自动执行）。"""

import argparse
import io
import json
import os
import secrets
import time
import uuid
from contextlib import ExitStack
from pathlib import Path

import httpx
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
SOURCE = """以下为虚构测试案例，仅用于验证软件，不代表真实项目业绩。

背景：校园失物信息分散在多个群聊，查找需要人工翻阅消息。

问题：发布记录缺少统一入口，物品分类与认领状态不明确。

方案：提供失物发布、分类筛选、关键词搜索和认领状态更新。未实现图片相似度搜索。

实现：前端采用 React，后端采用 FastAPI，数据保存在 PostgreSQL。
测试案例中的个人分工为页面开发、搜索接口与本地测试。

验证：在本地用12条虚构物品记录验证发布和关键词检索；其中8条记录标为待认领、4条标为已认领。没有真实用户测试，也没有效率提升比例。

不足：认领身份仍需人工确认。下一步计划补充异常输入测试和用户可用性测试，尚未实施。
"""


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:39800")
    parser.add_argument("--email", help="已有本地账号；密码通过 AIPPT_TEST_PASSWORD 环境变量传入")
    parser.add_argument("--upload-source", action="store_true", help="通过文件上传测试解析与存储")
    parser.add_argument("--page-count", type=int, choices=range(5, 21), default=6)
    parser.add_argument("--duration-minutes", type=int, choices=range(1, 61), default=5)
    parser.add_argument("--source-file", type=Path, help="自定义虚构验收材料（UTF-8）")
    parser.add_argument("--output-dir", type=Path, help="独立保存本次报告和导出文件")
    parser.add_argument("--model", help="仅覆盖验收账号的模型名称，需同时使用 --personal-model")
    parser.add_argument(
        "--audience-profile",
        choices=["academic", "executive", "technical", "investor", "learner", "custom"],
        help="启用听众叙事并验证实际规划",
    )
    parser.add_argument(
        "--personal-model",
        action="store_true",
        help="用服务器配置在验收账号中测试个人模型；不输出密钥",
    )
    args = parser.parse_args()
    if args.model and not args.personal_model:
        parser.error("--model requires --personal-model")
    source_text = args.source_file.read_text(encoding="utf-8") if args.source_file else SOURCE
    report = {"requested_pages": args.page_count, "duration_minutes": args.duration_minutes}
    out = ROOT / "backend/var/live-validation"
    if args.audience_profile:
        out = out / ("narrative-" + args.audience_profile)
    if args.output_dir:
        out = args.output_dir.resolve()
    elif args.page_count != 6:
        out = out / f"pages-{args.page_count}"
    out.mkdir(parents=True, exist_ok=True)
    with (
        httpx.Client(base_url=args.url + "/api/v1", timeout=90, trust_env=False) as client,
        ExitStack() as cleanup,
    ):

        def save_report():
            (out / "report.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
            )

        def record_outcome(exc_type, _error, _traceback):
            report["passed"] = exc_type is None and report.get("pptx_verified", False)
            if exc_type:
                report["failure_type"] = exc_type.__name__
            save_report()

        cleanup.push(record_outcome)

        def call(method, path, **kwargs):
            response = client.request(method, path, **kwargs)
            response.raise_for_status()
            return response

        def wait_for(path, ready, timeout):
            deadline = time.monotonic() + timeout
            previous = None
            while time.monotonic() < deadline:
                value = call("GET", path).json()
                status = (value["status"], value.get("ready"), value.get("failed"))
                if status != previous:
                    print(f"Progress: {status}", flush=True)
                    previous = status
                if value["status"] == "failed":
                    report["failed_endpoint"] = path.rsplit("/", 1)[-1]
                    report["failure_message"] = value.get("error")
                    raise RuntimeError(value.get("error", "Generation failed"))
                if ready(value):
                    return value
                if value["status"] == "partial" and value.get("failed", 0):
                    (out / "failed-deck.json").write_text(
                        json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                    raise RuntimeError("Some slides failed; inspect local failed-deck.json")
                time.sleep(2)
            raise TimeoutError("Generation did not complete within the test time limit")

        health = call("GET", "/health").json()
        if health.get("generation_mode") != "live" or not health.get("llm_configured"):
            raise RuntimeError("Restart API and worker in live mode with a configured key first")
        report["health"] = health
        if args.email:
            password = os.environ.get("AIPPT_TEST_PASSWORD")
            if not password:
                raise RuntimeError("AIPPT_TEST_PASSWORD is required with --email")
            auth = call("POST", "/auth/login", json={"email": args.email, "password": password})
        else:
            auth = call(
                "POST",
                "/auth/register",
                json={
                    "email": f"live-{uuid.uuid4().hex}@example.com",
                    "password": secrets.token_urlsafe(24),
                },
            )
        client.headers["Authorization"] = f"Bearer {auth.json()['access_token']}"
        if args.personal_model:
            from dotenv import dotenv_values

            config = dotenv_values(ROOT / "backend/.env")
            personal = {
                "base_url": config["LLM_BASE_URL"],
                "model": args.model or config["LLM_MODEL"],
                "api_key": config["LLM_API_KEY"],
            }
            saved = call("PUT", "/model-settings", json=personal)
            assert saved.json()["custom"] and personal["api_key"] not in saved.text

            def remove_personal_key():
                assert not call("DELETE", "/model-settings").json()["custom"]
                report["personal_key_removed"] = True

            cleanup.callback(remove_personal_key)
            report["personal_model"] = saved.json()["model"]
            probe = call(
                "POST",
                "/model-settings/test",
                json={k: v for k, v in personal.items() if k != "api_key"},
            )
            report["personal_probe"] = probe.json()["ok"]
            assert probe.json()["ok"], "Personal model connection probe failed"
            report["personal_probe"] = True
            print("Personal model saved; connection verified", flush=True)
        project = call(
            "POST",
            "/projects",
            json={
                "title": f"真实 API 验收 · {args.page_count}页校园失物招领（虚构案例）",
                "page_count": args.page_count,
                "theme_id": "clear" if args.audience_profile else "ivory",
                "audience": {
                    "executive": "负责试点决策的管理者",
                    "learner": "首次接触应用开发的听众",
                    "academic": "项目评审老师",
                }.get(args.audience_profile, "项目评审老师"),
                "content_density": "concise",
                "layout_mode": "flex",
                "brief": {
                    "scenario": "defense",
                    "duration_minutes": args.duration_minutes,
                    "focus": "个人贡献、已验证结果与尚未验证的边界",
                    "narrative_enabled": bool(args.audience_profile),
                    "audience_profile": args.audience_profile or "custom",
                    "knowledge_level": "newcomer"
                    if args.audience_profile == "learner"
                    else "familiar",
                    "narrative_overrides": {
                        "executive": (
                            "先给是否开展试点的建议与决策请求，少讲背景。明确哪些仍是待验证假设。"
                        ),
                        "learner": (
                            "面向首次接触应用开发的听众，"
                            "用熟悉场景解释检索与状态管理，最后留一个练习。"
                        ),
                        "academic": "面向评审老师，解释研究问题、方法取舍、验证证据和局限。",
                    }.get(args.audience_profile, ""),
                },
            },
        ).json()
        url = f"/projects/{project['id']}"
        report["project_id"] = project["id"]
        print(f"Project: {project['id']}", flush=True)
        if args.upload_source:
            source = call(
                "POST",
                url + "/sources/upload",
                files={"file": ("validation.txt", source_text.encode("utf-8"), "text/plain")},
            ).json()
        else:
            source = call(
                "POST", url + "/sources", json={"kind": "text", "content": source_text}
            ).json()
        report["source_kind"] = source["kind"]
        report["source_char_count"] = source["char_count"]
        assert source["char_count"] > 0
        start = time.monotonic()
        call("POST", url + "/outline/generate")
        outline = wait_for(
            url + "/outline",
            lambda value: value["status"] == "draft",
            # 覆盖两次各 600 秒的大纲任务预算及排队余量，不在任务仍运行时清除密钥。
            1320,
        )
        report["outline_seconds"] = round(time.monotonic() - start, 2)
        report["outline_pages"] = len(outline["pages"])
        assert len(outline["pages"]) == args.page_count
        if args.audience_profile:
            assert outline["narrative"]["source"] == "ai", "Expected a real AI narrative plan"
            assert (
                sum(page["speaker_seconds"] for page in outline["pages"])
                == args.duration_minutes * 60
            )
            assert all(page["narrative_role"] for page in outline["pages"])
            report["audience_profile"] = args.audience_profile
            report["narrative"] = outline["narrative"]
            report["page_titles"] = [page["title"] for page in outline["pages"]]
            report["time_budget_seconds"] = args.duration_minutes * 60
        (out / "outline.json").write_text(
            json.dumps(outline, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        call("POST", url + "/outline/confirm", json={"revision": outline["revision"]})
        start = time.monotonic()
        call("POST", url + "/deck/generate", json={})
        deck = wait_for(
            url + "/deck", lambda value: value["status"] == "ready", max(600, args.page_count * 60)
        )
        report["deck_seconds"] = round(time.monotonic() - start, 2)
        report["ready_pages"] = deck["ready"]
        report["failed_pages"] = deck["failed"]
        assert deck["ready"] == args.page_count and deck["failed"] == 0
        quality = call("GET", url + "/deck/quality").json()
        report["export_allowed"] = quality["export_allowed"]
        report["quality_issues"] = quality["issues"]
        (out / "deck.json").write_text(
            json.dumps(deck, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        (out / "report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        exported = call("GET", url + "/deck/export")
        presentation = Presentation(io.BytesIO(exported.content))
        report["exported_pages"] = len(presentation.slides)
        assert len(presentation.slides) == args.page_count
        assert all(
            any(shape.has_text_frame for shape in slide.shapes) for slide in presentation.slides
        )
        assert all(
            "大纲要点来源" in slide.notes_slide.notes_text_frame.text
            for slide in presentation.slides
        )
        output = ROOT / "examples/live-api-defense.pptx"
        if args.audience_profile:
            output = ROOT / f"examples/narrative-{args.audience_profile}.pptx"
            assert all(
                "讲述提示" in slide.notes_slide.notes_text_frame.text
                for slide in presentation.slides
            )
        if args.output_dir or args.page_count != 6:
            output = out / "presentation.pptx"
        output.write_bytes(exported.content)
        report["pptx_bytes"] = len(exported.content)
        report["pptx_verified"] = True
        (out / "report.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        print(
            json.dumps(
                {k: v for k, v in report.items() if k not in ("quality_issues", "narrative")},
                ensure_ascii=False,
            ),
            flush=True,
        )


if __name__ == "__main__":
    main()
